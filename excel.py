from    openpyxl            import load_workbook
from    openpyxl.styles     import Font
from    pathlib             import Path
import  datetime

from    data       import cell_address

def create_sheet(wb, class_num, title, class_list):
    template    = wb[str(class_num)]
    ws          = wb.copy_worksheet(template)
    ws.title    = title
    for i in range(class_num):
            ws[cell_address[class_num][i]]  = class_list[i]
            cell        = ws[cell_address[class_num][i]]
            cell.font   = Font(size=20, bold=True)
    ws["F1"]    = title
    return ws

def create_excel_file(soccer_list,       soccer_num,
            basketball_list,    basketball_num,
            volleyball_list,    volleyball_num,
            dodgeball_list,     dodgeball_num):
    
    wb=load_workbook("tornament_standard.xlsx")

    create_sheet(wb, soccer_num, "サッカー", soccer_list)

    create_sheet(wb, basketball_num, "バスケ", basketball_list)

    create_sheet(wb, volleyball_num, "バレー", volleyball_list)

    create_sheet(wb, dodgeball_num, "ドッヂ", dodgeball_list)

    keep_sheets = {"サッカー", "バスケ", "バレー", "ドッヂ", "timetable"}

    for sheet_name in wb.sheetnames[:]:
        if sheet_name not in keep_sheets:
            del wb[sheet_name]

    
    now         = datetime.datetime.now()
    save_dir    = Path("ExcelFile")
    save_dir.mkdir(parents=True, exist_ok=True)
    filename    = now.strftime("%y%m%d_%H%M%S.xlsx")
    save_path   = save_dir / filename
    wb.save(save_path)
